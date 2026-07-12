import base64
import hashlib
import io
import json
import math
import time
from html.parser import HTMLParser
from pathlib import Path
from urllib import error, request

try:
    from app_database import save_document_metadata
    from graph_rag import graph_store_status, retrieve_graph_context, upsert_graph_chunks
except ModuleNotFoundError:
    from backend.app_database import save_document_metadata
    from backend.graph_rag import graph_store_status, retrieve_graph_context, upsert_graph_chunks


class HtmlTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self._chunks = []

    def handle_data(self, data):
        self._chunks.append(data)

    def get_text(self):
        return "".join(self._chunks)


def html_to_text(html):
    parser = HtmlTextExtractor()
    parser.feed(html)
    return parser.get_text()


def extract_text_from_pdf(file_bytes):
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            return "\n\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception:
        pass

    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        raise ValueError(
            "PDF ingestion requires pdfplumber or PyPDF2. Install one with `pip install pdfplumber PyPDF2`."
        )


def extract_text_from_docx(file_bytes):
    try:
        import docx

        document = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        tables = []
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        tables.append(cell.text)
        return "\n".join(paragraphs + tables)
    except Exception:
        raise ValueError(
            "DOCX ingestion requires python-docx. Install it with `pip install python-docx`."
        )


def extract_text_from_pptx(file_bytes):
    try:
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(file_bytes))
        texts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    texts.append(shape.text)
                if shape.shape_type == 19 and shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text:
                                texts.append(cell.text)
        return "\n".join(texts)
    except Exception:
        raise ValueError(
            "PPTX ingestion requires python-pptx. Install it with `pip install python-pptx`."
        )


def extract_text_from_bytes(file_bytes, file_name, mime_type=None):
    extension = Path(file_name).suffix.lower()
    if extension in {".txt", ".md", ".csv", ".json", ".py", ".js", ".css"}:
        return file_bytes.decode("utf-8", errors="replace")
    if extension in {".html", ".htm"}:
        return html_to_text(file_bytes.decode("utf-8", errors="replace"))
    if extension == ".pdf":
        return extract_text_from_pdf(file_bytes)
    if extension == ".docx":
        return extract_text_from_docx(file_bytes)
    if extension == ".pptx":
        return extract_text_from_pptx(file_bytes)
    if extension == ".ppt":
        raise ValueError(
            "Legacy PPT (.ppt) is not supported. Please convert to PPTX and try again."
        )
    if mime_type and mime_type.startswith("text/"):
        return file_bytes.decode("utf-8", errors="replace")
    raise ValueError(
        f"Unsupported document type '{extension}'. Supported file types include PDF, DOCX, PPTX, HTML, TXT, MD, CSV, JSON, PY, JS, and CSS."
    )


def decode_document_payload(doc):
    if doc.get("text") is not None:
        return doc["text"]
    if doc.get("fileData"):
        file_name = doc.get("fileName") or doc.get("title") or "document"
        mime_type = doc.get("mimeType")
        file_bytes = base64.b64decode(doc["fileData"])
        return extract_text_from_bytes(file_bytes, file_name, mime_type)
    raise ValueError("Document payload must include either text or encoded fileData for ingestion." )


class HtmlTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self._chunks = []

    def handle_data(self, data):
        self._chunks.append(data)

    def get_text(self):
        return "".join(self._chunks)


def html_to_text(html):
    parser = HtmlTextExtractor()
    parser.feed(html)
    return parser.get_text()


def extract_text_from_pdf(file_bytes):
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            return "\n\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception:
        pass

    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        raise ValueError(
            "PDF ingestion requires pdfplumber or PyPDF2. Install one with `pip install pdfplumber PyPDF2`."
        )


def extract_text_from_docx(file_bytes):
    try:
        import docx

        document = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        tables = []
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        tables.append(cell.text)
        return "\n".join(paragraphs + tables)
    except Exception:
        raise ValueError(
            "DOCX ingestion requires python-docx. Install it with `pip install python-docx`."
        )


def extract_text_from_pptx(file_bytes):
    try:
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(file_bytes))
        texts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    texts.append(shape.text)
                if shape.shape_type == 19 and shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text:
                                texts.append(cell.text)
        return "\n".join(texts)
    except Exception:
        raise ValueError(
            "PPTX ingestion requires python-pptx. Install it with `pip install python-pptx`."
        )


def extract_text_from_bytes(file_bytes, file_name, mime_type=None):
    extension = Path(file_name).suffix.lower()
    if extension in {".txt", ".md", ".csv", ".json", ".py", ".js", ".css"}:
        return file_bytes.decode("utf-8", errors="replace")
    if extension in {".html", ".htm"}:
        return html_to_text(file_bytes.decode("utf-8", errors="replace"))
    if extension == ".pdf":
        return extract_text_from_pdf(file_bytes)
    if extension == ".docx":
        return extract_text_from_docx(file_bytes)
    if extension == ".pptx":
        return extract_text_from_pptx(file_bytes)
    if extension == ".ppt":
        raise ValueError(
            "Legacy PPT (.ppt) is not supported. Please convert to PPTX and try again."
        )
    if mime_type and mime_type.startswith("text/"):
        return file_bytes.decode("utf-8", errors="replace")
    raise ValueError(
        f"Unsupported document type '{extension}'. Supported file types include PDF, DOCX, PPTX, HTML, TXT, MD, CSV, JSON, PY, JS, and CSS."
    )


def decode_document_payload(doc):
    if doc.get("text") is not None:
        return doc["text"]
    if doc.get("fileData"):
        file_name = doc.get("fileName") or doc.get("title") or "document"
        mime_type = doc.get("mimeType")
        file_bytes = base64.b64decode(doc["fileData"])
        return extract_text_from_bytes(file_bytes, file_name, mime_type)
    raise ValueError("Document payload must include either text or encoded fileData for ingestion." )


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
CHROMA_DIR = DATA_DIR / "chroma"
FAISS_DIR = DATA_DIR / "faiss"
FALLBACK_VECTOR_FILE = DATA_DIR / "vector_store.json"
DEFAULT_COLLECTION = "course_docs"
DEFAULT_EMBEDDING_MODEL = "nomic-embed-text"
DEFAULT_BASE_URL = "http://127.0.0.1:11434"
DEFAULT_VECTOR_BACKEND = "auto"
LOCAL_JSON_BACKEND = "local-json-fallback"
SUPPORTED_VECTOR_BACKENDS = {
    "auto": "Auto (Chroma if installed, otherwise local JSON)",
    "chroma": "Chroma",
    LOCAL_JSON_BACKEND: "Local JSON fallback",
    "faiss": "FAISS",
}


def vector_store_status():
    return {
        "backend": resolve_vector_backend(DEFAULT_VECTOR_BACKEND),
        "defaultBackend": DEFAULT_VECTOR_BACKEND,
        "supportedBackends": [
            backend_status(name, label)
            for name, label in SUPPORTED_VECTOR_BACKENDS.items()
        ],
        "chromaInstalled": chroma_available(),
        "faissInstalled": faiss_available(),
        "chromaPath": str(CHROMA_DIR),
        "faissPath": str(FAISS_DIR),
        "fallbackPath": str(FALLBACK_VECTOR_FILE),
    }


def chroma_available():
    try:
        import chromadb  # noqa: F401

        return True
    except Exception:
        return False


def faiss_available():
    try:
        import faiss  # noqa: F401

        return True
    except Exception:
        return False


def normalize_vector_backend(value):
    backend = (value or DEFAULT_VECTOR_BACKEND).strip().lower()
    aliases = {
        "json": LOCAL_JSON_BACKEND,
        "local": LOCAL_JSON_BACKEND,
        "local-json": LOCAL_JSON_BACKEND,
        "fallback": LOCAL_JSON_BACKEND,
    }
    backend = aliases.get(backend, backend)
    if backend not in SUPPORTED_VECTOR_BACKENDS:
        raise ValueError(
            f"Unsupported vector backend '{value}'. "
            f"Choose one of: {', '.join(SUPPORTED_VECTOR_BACKENDS)}."
        )
    return backend


def resolve_vector_backend(requested=None):
    backend = normalize_vector_backend(requested)
    if backend == "auto":
        return "chroma" if chroma_available() else LOCAL_JSON_BACKEND
    return backend


def backend_status(name, label):
    if name == "auto":
        resolved = resolve_vector_backend(name)
        return {
            "name": name,
            "label": label,
            "available": True,
            "resolvedBackend": resolved,
        }
    if name == "chroma":
        return {
            "name": name,
            "label": label,
            "available": chroma_available(),
            "message": "Install chromadb to use this backend." if not chroma_available() else "",
        }
    if name == "faiss":
        return {
            "name": name,
            "label": label,
            "available": False,
            "installed": faiss_available(),
            "message": "Adapter scaffolded; persistence/retrieval implementation is pending.",
        }
    return {
        "name": name,
        "label": label,
        "available": True,
        "message": "Stores vectors in data/vector_store.json.",
    }


def require_backend_available(backend):
    if backend == "chroma" and not chroma_available():
        raise ValueError("Chroma backend selected, but chromadb is not installed.")
    if backend == "faiss":
        raise ValueError(
            "FAISS backend is selectable as an adapter scaffold, but ingestion/retrieval "
            "is not implemented yet. Use Chroma or local JSON fallback for working runs."
        )


def list_vector_collections():
    collections = {}
    errors = []

    if chroma_available():
        try:
            import chromadb

            client = chromadb.PersistentClient(path=str(CHROMA_DIR))
            for collection in client.list_collections():
                name = collection.name if hasattr(collection, "name") else str(collection)
                count = None
                try:
                    count = client.get_collection(name).count()
                except Exception:
                    pass
                collections[name] = {
                    "name": name,
                    "vectorCount": count,
                    "backend": "chroma",
                    "backends": ["chroma"],
                }
        except Exception as exc:
            errors.append(f"Chroma: {exc}")

    fallback_store = load_fallback_store()
    for name, items in fallback_store.get("collections", {}).items():
        existing = collections.setdefault(
            name,
            {
                "name": name,
                "vectorCount": 0,
                "backend": LOCAL_JSON_BACKEND,
                "backends": [],
            },
        )
        existing.setdefault("backends", [])
        if LOCAL_JSON_BACKEND not in existing["backends"]:
            existing["backends"].append(LOCAL_JSON_BACKEND)
        existing["fallbackCount"] = len(items)
        if existing.get("vectorCount") is None:
            existing["vectorCount"] = len(items)
        if existing.get("backend") != LOCAL_JSON_BACKEND:
            existing["backend"] = ", ".join(existing["backends"])

    return {
        "store": vector_store_status(),
        "collections": [
            value
            for key, value in sorted(collections.items())
        ],
        "errors": errors,
    }


def check_ollama_status(base_url=None):
    endpoint = (base_url or DEFAULT_BASE_URL).rstrip("/")
    try:
        req = request.Request(f"{endpoint}/api/tags", method="GET")
        with request.urlopen(req, timeout=2) as response:
            payload = json.loads(response.read().decode("utf-8"))
        models = [
            item.get("name")
            for item in payload.get("models", [])
            if item.get("name")
        ]
        return {
            "available": True,
            "baseUrl": endpoint,
            "models": models,
        }
    except Exception as exc:
        return {
            "available": False,
            "baseUrl": endpoint,
            "message": str(exc),
        }


def ingest_documents(payload):
    collection = payload.get("collection") or DEFAULT_COLLECTION
    embedding_model = payload.get("embeddingModel") or DEFAULT_EMBEDDING_MODEL
    base_url = payload.get("baseUrl") or DEFAULT_BASE_URL
    vector_backend = resolve_vector_backend(payload.get("vectorBackend"))
    require_backend_available(vector_backend)
    chunk_size = int(payload.get("chunkSize") or 700)
    chunk_overlap = int(payload.get("chunkOverlap") or 120)
    documents = payload.get("documents") or []

    if not documents and payload.get("text"):
        documents = [
            {
                "title": payload.get("title") or "Untitled document",
                "source": payload.get("source") or "manual",
                "text": payload.get("text"),
            }
        ]

    if not documents:
        raise ValueError("No documents were provided for ingestion.")

    chunks = []
    for doc in documents:
        text = decode_document_payload(doc)
        title = doc.get("title") or "Untitled document"
        source = doc.get("source") or "manual"
        doc_chunks = chunk_text(text, chunk_size, chunk_overlap)
        for index, chunk in enumerate(doc_chunks):
            chunks.append(
                {
                    "id": stable_id(collection, source, title, index, chunk),
                    "text": chunk,
                    "metadata": {
                        "title": title,
                        "source": source,
                        "chunkIndex": index,
                        "createdAt": time.time(),
                    },
                }
            )
        save_document_metadata(collection, title, source, len(doc_chunks))

    texts = [chunk["text"] for chunk in chunks]
    embeddings, embedding_backend = embed_texts(texts, embedding_model, base_url)

    if vector_backend == "chroma":
        upsert_chroma(collection, chunks, embeddings)
    elif vector_backend == "faiss":
        upsert_faiss(collection, chunks, embeddings)
    else:
        upsert_fallback(collection, chunks, embeddings)

    graph_result = {"status": "disabled"}
    if payload.get("graphEnabled", True):
        graph_result = upsert_graph_chunks(collection, chunks, payload)

    return {
        "status": "ingested",
        "collection": collection,
        "chunks": len(chunks),
        "vectorBackend": vector_backend,
        "embeddingBackend": embedding_backend,
        "graph": graph_result,
        "store": vector_store_status(),
    }


def retrieve_context(config, query):
    collection = config.get("collection") or DEFAULT_COLLECTION
    embedding_model = config.get("embeddingModel") or DEFAULT_EMBEDDING_MODEL
    base_url = config.get("baseUrl") or DEFAULT_BASE_URL
    vector_backend = resolve_vector_backend(config.get("vectorBackend"))
    top_k = int(config.get("topK") or 3)
    retrieval_mode = (config.get("retrievalMode") or "vector").lower()
    if retrieval_mode not in {"vector", "graph", "hybrid"}:
        retrieval_mode = "vector"
    if retrieval_mode != "graph":
        require_backend_available(vector_backend)
    graph_top_k = int(config.get("graphTopK") or top_k)
    graph_hops = int(config.get("graphHops") or 1)
    graph_answer_mode = (config.get("graphAnswerMode") or "grounded").lower()
    if graph_answer_mode not in {"grounded", "hybrid"}:
        graph_answer_mode = "grounded"

    if not query.strip():
        return {
            "context": "Retriever received an empty query.",
            "matches": [],
            "vectorBackend": vector_backend,
            "retrievalMode": retrieval_mode,
        }

    if retrieval_mode == "graph":
        graph_retrieval = retrieve_graph_context(collection, query, graph_top_k, graph_hops, config)
        graph_evidence = dict(graph_retrieval.get("graphEvidence") or {})
        graph_evidence.setdefault("answerMode", graph_retrieval.get("graphAnswerMode") or graph_answer_mode)
        graph_evidence.setdefault("fallbackAllowed", graph_evidence.get("answerMode") == "hybrid")
        graph_evidence.setdefault("fallbackUsed", False)
        return {
            **graph_retrieval,
            "graphEvidence": graph_evidence,
            "retrievalMode": "graph",
            "graphAnswerMode": graph_evidence.get("answerMode"),
            "vectorBackend": vector_backend,
            "embeddingBackend": None,
        }

    vector_retrieval = retrieve_vector_context(
        collection,
        query,
        embedding_model,
        base_url,
        vector_backend,
        top_k,
    )
    if retrieval_mode == "hybrid":
        graph_retrieval = retrieve_graph_context(collection, query, graph_top_k, graph_hops, config)
        combined_matches = merge_retrieval_matches(
            vector_retrieval.get("matches", []),
            graph_retrieval.get("matches", []),
        )
        graph_evidence = dict(graph_retrieval.get("graphEvidence") or {})
        graph_evidence.setdefault("answerMode", graph_retrieval.get("graphAnswerMode") or graph_answer_mode)
        graph_evidence.setdefault("fallbackAllowed", graph_evidence.get("answerMode") == "hybrid")
        graph_evidence.setdefault("fallbackUsed", False)
        graph_has_evidence = bool(
            graph_evidence.get("paths")
            or graph_evidence.get("relationships")
            or graph_evidence.get("entities")
        )
        vector_context = vector_retrieval.get("context")
        if graph_has_evidence and not vector_retrieval.get("matches"):
            vector_context = ""
        context_parts = [
            part for part in [
                vector_context,
                graph_retrieval.get("context"),
            ]
            if part
        ]
        return {
            "context": "\n\n---\n\n".join(context_parts) or "No hybrid retrieval context found.",
            "matches": combined_matches,
            "vectorBackend": vector_backend,
            "embeddingBackend": vector_retrieval.get("embeddingBackend"),
            "retrievalMode": "hybrid",
            "graphAnswerMode": graph_retrieval.get("graphAnswerMode") or graph_answer_mode,
            "graphBackend": graph_retrieval.get("graphBackend"),
            "graphEvidence": graph_evidence,
        }

    return {**vector_retrieval, "retrievalMode": "vector"}


def retrieve_vector_context(collection, query, embedding_model, base_url, vector_backend, top_k):
    query_embedding, embedding_backend = embed_texts([query], embedding_model, base_url)
    if vector_backend == "chroma":
        matches = query_chroma(collection, query_embedding[0], top_k)
    elif vector_backend == "faiss":
        matches = query_faiss(collection, query_embedding[0], top_k)
    else:
        matches = query_fallback(collection, query_embedding[0], top_k)

    if not matches:
        return {
            "context": (
                f"No retrieved context found in collection '{collection}'. "
                "Ingest documents first, then run the workflow again."
            ),
            "matches": [],
            "vectorBackend": vector_backend,
            "embeddingBackend": embedding_backend,
        }

    lines = [
        f"Retrieved context from collection '{collection}' using {vector_backend}:",
    ]
    for index, match in enumerate(matches, start=1):
        title = match["metadata"].get("title", "Untitled")
        score = round(match.get("score", 0), 4)
        lines.append(f"[{index}] {title} (score: {score})\n{match['text']}")

    return {
        "context": "\n\n".join(lines),
        "matches": matches,
        "vectorBackend": vector_backend,
        "embeddingBackend": embedding_backend,
    }


def merge_retrieval_matches(vector_matches, graph_matches):
    merged = []
    seen = set()
    for stage, matches in (("vector", vector_matches), ("graph", graph_matches)):
        for match in matches:
            metadata = dict(match.get("metadata", {}))
            match_id = metadata.get("id") or stable_id(stage, metadata.get("source"), metadata.get("title"), match.get("text"))
            if match_id in seen:
                continue
            seen.add(match_id)
            metadata.setdefault("stage", stage)
            merged.append({**match, "metadata": metadata})
    return merged


def chunk_text(text, chunk_size, overlap):
    cleaned = " ".join(text.split())
    if not cleaned:
        return []

    chunks = []
    start = 0
    while start < len(cleaned):
        end = min(start + chunk_size, len(cleaned))
        chunks.append(cleaned[start:end])
        if end == len(cleaned):
            break
        start = max(0, end - overlap)
    return chunks


def embed_texts(texts, model, base_url):
    try:
        payload = {"model": model, "input": texts}
        data = json.dumps(payload).encode("utf-8")
        req = request.Request(
            f"{base_url.rstrip('/')}/api/embed",
            data=data,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with request.urlopen(req, timeout=90) as response:
            result = json.loads(response.read().decode("utf-8"))
        embeddings = result.get("embeddings")
        if embeddings and len(embeddings) == len(texts):
            return embeddings, f"ollama:{model}"
    except error.URLError:
        pass
    except Exception:
        pass

    return [hash_embedding(text) for text in texts], "local-hash-fallback"


def hash_embedding(text, dimensions=64):
    vector = [0.0] * dimensions
    for token in text.lower().split():
        digest = hashlib.sha256(token.encode("utf-8")).digest()
        index = int.from_bytes(digest[:2], "big") % dimensions
        sign = 1.0 if digest[2] % 2 == 0 else -1.0
        vector[index] += sign
    return normalize(vector)


def normalize(vector):
    magnitude = math.sqrt(sum(value * value for value in vector))
    if magnitude == 0:
        return vector
    return [value / magnitude for value in vector]


def cosine_similarity(left, right):
    return sum(a * b for a, b in zip(left, right))


def stable_id(*parts):
    text = "::".join(str(part) for part in parts)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:24]


def upsert_chroma(collection, chunks, embeddings):
    import chromadb

    CHROMA_DIR.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    chroma_collection = client.get_or_create_collection(name=collection)
    chroma_collection.upsert(
        ids=[chunk["id"] for chunk in chunks],
        documents=[chunk["text"] for chunk in chunks],
        embeddings=embeddings,
        metadatas=[chunk["metadata"] for chunk in chunks],
    )


def query_chroma(collection, query_embedding, top_k):
    import chromadb

    CHROMA_DIR.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    chroma_collection = client.get_or_create_collection(name=collection)
    result = chroma_collection.query(
        query_embeddings=[query_embedding],
        n_results=top_k,
        include=["documents", "metadatas", "distances"],
    )

    documents = result.get("documents", [[]])[0]
    metadatas = result.get("metadatas", [[]])[0]
    distances = result.get("distances", [[]])[0]
    matches = []
    for doc, metadata, distance in zip(documents, metadatas, distances):
        matches.append({
            "text": doc,
            "metadata": metadata or {},
            "score": 1 / (1 + float(distance)),
        })
    return matches


def upsert_faiss(collection, chunks, embeddings):
    raise ValueError(
        "FAISS adapter is scaffolded but persistence is not implemented yet. "
        "Use Chroma or local JSON fallback for working ingestion."
    )


def query_faiss(collection, query_embedding, top_k):
    raise ValueError(
        "FAISS adapter is scaffolded but retrieval is not implemented yet. "
        "Use Chroma or local JSON fallback for working retrieval."
    )


def load_fallback_store():
    if not FALLBACK_VECTOR_FILE.exists():
        return {"collections": {}}
    return json.loads(FALLBACK_VECTOR_FILE.read_text(encoding="utf-8"))


def save_fallback_store(store):
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    FALLBACK_VECTOR_FILE.write_text(json.dumps(store, indent=2), encoding="utf-8")


def upsert_fallback(collection, chunks, embeddings):
    store = load_fallback_store()
    existing = {
        item["id"]: item
        for item in store.setdefault("collections", {}).setdefault(collection, [])
    }
    for chunk, embedding in zip(chunks, embeddings):
        existing[chunk["id"]] = {
            "id": chunk["id"],
            "text": chunk["text"],
            "metadata": chunk["metadata"],
            "embedding": embedding,
        }
    store["collections"][collection] = list(existing.values())
    save_fallback_store(store)


def query_fallback(collection, query_embedding, top_k):
    store = load_fallback_store()
    items = store.get("collections", {}).get(collection, [])
    scored = []
    for item in items:
        scored.append({
            "text": item["text"],
            "metadata": item.get("metadata", {}),
            "score": cosine_similarity(query_embedding, item.get("embedding", [])),
        })
    return sorted(scored, key=lambda item: item["score"], reverse=True)[:top_k]
